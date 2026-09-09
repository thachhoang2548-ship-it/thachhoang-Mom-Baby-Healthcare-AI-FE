import os
import glob

def find_and_read():
    import docx
    target = os.path.join(chr(68) + ":", os.sep, "Downloads", "SDS_MomBaby_Healthcare_AI_v4 (2).docx")
    if not os.path.exists(target):
        print("Not found at", target)
        return
    print(f"Reading: {target}")
    doc = docx.Document(target)
    
    headings = []
    text = []
    for p in doc.paragraphs:
        if p.text.strip():
            if p.style.name.startswith("Heading"):
                headings.append(f"{p.style.name}: {p.text.strip()}")
            text.append(p.text.strip())
            
    print("\n--- HEADINGS ---")
    for h in headings:
        print(h)
        
    print("\n--- SAMPLE PARAGRAPHS ---")
    for t in text[:50]:
        print("-", t)

    print("\n--- TABLES SUMMARY ---")
    print(f"Total tables: {len(doc.tables)}")
    for i, t in enumerate(doc.tables[:15]):
        print(f"\nTable {i+1} rows: {len(t.rows)}")
        for r in t.rows[:6]:
            row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells if c.text.strip()]
            if row_vals:
                print("  | ", " | ".join(row_vals))

if __name__ == "__main__":
    find_and_read()
