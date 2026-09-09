import os
import docx

target = "SDS_MomBaby_Healthcare_AI.docx"
out_file = "sds_summary_v7.txt"

if os.path.exists(target):
    doc = docx.Document(target)
    lines = []
    lines.append("=== PARAGRAPHS ===")
    for p in doc.paragraphs:
        if p.text.strip():
            style = p.style.name if p.style else ""
            lines.append(f"[{style}] {p.text.strip()}")
            
    lines.append("\n=== TABLES ===")
    for i, t in enumerate(doc.tables):
        lines.append(f"\n--- Table {i+1} ---")
        for r in t.rows:
            row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells if c.text.strip()]
            if row_vals:
                lines.append(" | ".join(row_vals))
                
    with open(out_file, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Successfully wrote {len(lines)} lines to {out_file} from {target}")
else:
    print(f"File not found: {target}")
